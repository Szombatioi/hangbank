# Usage: 
# $ python .\main.py test.pdf -o E:\Work\test_python_corpus_maker

import nltk
nltk.download('punkt')
nltk.download('punkt_tab')

import re
from nltk.tokenize import sent_tokenize
import os
import sys
from docx import Document
from PyPDF2 import PdfReader
from datetime import datetime

def extract_text_from_txt(path):
    with open(path, 'r', encoding='utf-8') as file:
        return file.read()

def extract_text_from_docx(file_path):
    doc = Document(file_path)
    return '\n'.join([p.text for p in doc.paragraphs])

def extract_text_from_pdf(file_path):
    reader = PdfReader(file_path)
    all_text = []

    for page in reader.pages:
        page_text = page.extract_text() or ""

        # # Remove hyphenation at line breaks (like "adipiscing-\nelit")
        # page_text = re.sub(r'-\n\s*', '', page_text)

        # # Replace single newlines inside paragraphs with space
        # page_text = re.sub(r'(?<!\n)\n(?!\n)', ' ', page_text)

        # # Remove leading bullets or dashes at the start of sentences
        # page_text = re.sub(r'(?<=\n)-\s+', '', page_text)  # after a newline
        # page_text = re.sub(r'^-\s+', '', page_text)        # start of text

        # # Split into sentences
        # sentences = sent_tokenize(page_text)

        # # Remove empty sentences
        # sentences = [re.sub(r'\s+', ' ', s.strip()) for s in sentences if s.strip()]

        # all_text.append("\n".join(sentences))
        all_text.append(page_text)
    
    return [t for t in all_text if len(t) > 0]

def clean_text(text):
    abbreviations = [
        "Mr", "Mrs", "Ms", "Dr", "Prof", "Sr", "Jr", "vs", "etc",
        "e.g", "i.e", "Fig", "fig", "Eq", "eq", "Inc", "Ltd", "St"
    ]

    # 1. Fix hyphenation at line breaks (e.g. "adipis-\ncing" -> "adipiscing")
    text = re.sub(r'-\s*\n\s*', '', text)

    # 2. Protect newlines that indicate new sentences (like bullet points or dashes)
    # Replace them with a placeholder so we don’t accidentally merge them
    text = re.sub(r'\n\s*[-•]\s*', ' <NEWLINE_DASH> ', text)

    # 3. Merge remaining newlines (just noise from PDF)
    text = re.sub(r'\s*\n\s*', ' ', text)

    # 4. Normalize multiple spaces
    text = re.sub(r'\s+', ' ', text).strip()

    # 5. Temporarily protect abbreviations ("Dr." -> "Dr<ABBR>")
    for abbr in abbreviations:
        text = re.sub(rf'\b{abbr}\.', f'{abbr}<ABBR>', text)

    # 6. Split by sentence-ending punctuation (., !, ?)
    sentences = re.split(r'(?<=[.!?])\s+(?=[A-Z0-9\-])', text)

    # 7. Restore abbreviations and placeholders
    restored = []
    for s in sentences:
        s = s.replace('<ABBR>', '.')
        s = re.sub(r'\s*<NEWLINE_DASH>\s*', '\n- ', s)  # keep as visible dash line
        restored.append(s.strip())

    # 8. Optionally split on real newlines that survived (like dash lines)
    final_sentences = []
    for s in restored:
        parts = [p.strip() for p in re.split(r'\n+', s) if p.strip()]
        final_sentences.extend(parts)

    return "\n".join(final_sentences)

    # sentences = "\n".join(sent_tokenize(text))
    # sentences = re.sub(r'(?<=\n)-\s+', '', sentences)
    # sentences = re.sub(r'^-\s+', '', sentences)
    # return sentences

    # # Remove hyphenation at line breaks (like "adipiscing-\nelit")
    # text = re.sub(r'-\n\s*', '', text)

    # # Replace single newlines inside paragraphs with space
    # text = re.sub(r'(?<!\n)\n(?!\n)', ' ', text)

    # # Remove leading bullets or dashes at the start of sentences
    # text = re.sub(r'(?<=\n)-\s+', '', text)  # after a newline
    # text = re.sub(r'^-\s+', '', text)        # start of text

    # # Split into sentences
    # sentences = sent_tokenize(text)

    # # Remove empty sentences
    # sentences = [re.sub(r'\s+', ' ', s.strip()) for s in sentences if s.strip()]
    # return "\n".join(sentences)

# def convert_text_to_corpus(text):
#     for i, line in enumerate(text.split('\n')):
#         if len(line) == 0:
#             print("EMPTY LINE!")
#         else:
#             print(f"Block: {i}")
#             sentences = sent_tokenize(line)
#             sentences = [s.lstrip("- ") for s in sentences]
#             print("\n\t- ".join(sentences))
#         print("\n--------------------\n")

def convert_file_to_corpus(file_path, path=None):
    ext = os.path.splitext(file_path)[1].lower()
    text = None
    if ext == '.txt':
        text = extract_text_from_txt(file_path)
    elif ext == '.docx':
        text = extract_text_from_docx(file_path)
    elif ext == '.pdf':
        text = "".join(extract_text_from_pdf(file_path))
    else:
        raise ValueError("Unsupported file format. Please use .txt, .docx, or .pdf")
    cleaned_text = clean_text(text) #Cleaning text from \n and other artifacts
    
    # Saving cleaned text splitted up into separate text files
    total_lines = len(cleaned_text.split('\n'))
    # print(total_lines)
    for i, line in enumerate(cleaned_text.split('\n')):
        filename = re.sub(r'\.[^.]+$', '', file_path) + f"_{datetime.now().strftime("%Y%m%d%H%M%S")}" + f"_{str(i+1).zfill(len(str(total_lines)))}.txt"
        if(path is not None):
            if not os.path.exists(path):
                os.makedirs(path)
            filename = os.path.join(path, os.path.basename(filename))
        with open(filename, "w+") as file:
            file.write(line)
        print(filename)

if sys.argv.__len__() < 2:
    print("Error! Provide a file!")
    sys.exit(404)
args = sys.argv
path = None
if args.__contains__('-o') and len(args) >= args.index('-o') + 2:
    path = args[args.index('-o') + 1]

convert_file_to_corpus(sys.argv[1], path)